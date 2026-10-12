"""Build reports 02–08 from Markdown using the supplied academic report template."""
from pathlib import Path
import argparse
import hashlib
import json
import re
import shutil
import subprocess
import textwrap
import zipfile

import yaml
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
BASE = "6cfe25f87e4bac60550556fbb2b76c9931658b7e"
SECTIONS = json.loads((ROOT / "releases/report_sections.json").read_text())
TEMPLATE = ROOT / "templates/academic-report"
TICK = chr(96)
FENCE = TICK * 3
LINKS = {
    "gnu-modes": ("GNU Project", "Structure of File Mode Bits", "https://www.gnu.org/software/coreutils/manual/html_node/Mode-Structure.html"),
    "gnu-permissions": ("GNU Project", "Setting Permissions", "https://www.gnu.org/software/coreutils/manual/html_node/Setting-Permissions.html"),
    "gnu-chmod": ("GNU Project", "chmod Change access permissions", "https://www.gnu.org/software/coreutils/manual/html_node/chmod-invocation.html"),
    "linux-credentials": ("Linux man-pages project", "credentials(7)", "https://man7.org/linux/man-pages/man7/credentials.7.html"),
    "e2fs-chattr": ("e2fsprogs project", "chattr(1)", "https://man7.org/linux/man-pages/man1/chattr.1.html"),
    "linux-flags": ("Linux man-pages project", "FS_IOC_SETFLAGS(2)", "https://man7.org/linux/man-pages/man2/FS_IOC_GETFLAGS.2const.html"),
    "linux-execve": ("Linux man-pages project", "execve(2)", "https://man7.org/linux/man-pages/man2/execve.2.html"),
    "redhat-selinux": ("Red Hat", "SELinux User's and Administrator's Guide", "https://docs.redhat.com/en/documentation/red_hat_enterprise_linux/7/html/selinux_users_and_administrators_guide/index"),
    "redhat-labels": ("Red Hat", "Configuring SELinux for applications and services with non-standard configurations", "https://docs.redhat.com/en/documentation/red_hat_enterprise_linux/9/html/using_selinux/configuring-selinux-for-applications-and-services-with-non-standard-configurations_using-selinux"),
    "redhat-ports": ("Red Hat", "Changing port numbers", "https://docs.redhat.com/en/documentation/red_hat_enterprise_linux/6/html/managing_confined_services/sect-managing_confined_services-configuration_examples-changing_port_numbers"),
    "python-expressions": ("Python Software Foundation", "Expressions", "https://docs.python.org/3/reference/expressions.html"),
    "python-secrets": ("Python Software Foundation", "secrets Generate secure random numbers", "https://docs.python.org/3/library/secrets.html"),
}


def original(lab):
    path = f"labs/lab{lab:02}/report/infosec-intro--lab{lab:02}--report.md"
    return subprocess.check_output(["git", "show", f"{BASE}:{path}"], cwd=ROOT).decode()


def clean_heading(value):
    value = value.strip()
    value = value[0].upper() + value[1:]
    return value.replace(",", "").replace("=", " ").replace("+", " ").replace("/", " ")


def figure_steps(lab, source):
    if lab == 6:
        steps = [
            ("Запуск веб сервера", "Команда service httpd status показывает active running. В браузере открыта тестовая страница Apache. Состояние службы и HTTP-страница представлены на @fig-001.", "01.png", "fig-001"),
            ("Контекст процессов Apache", "Вывод ps aux -Z показывает контекст system_u:system_r:httpd_t:s0 для процессов веб-сервера. Тип процесса — httpd_t. См. @fig-002.", "02.png", "fig-002"),
            ("Переключатели SELinux для Apache", "Команда sestatus -b | grep httpd выводит переключатели политики. На снимке большинство значений off; httpd_builtin_scripting и httpd_enable_cgi имеют значение on. См. @fig-003.", "03.png", "fig-003"),
            ("Создание и чтение веб контента", "В /var/www/html создан test.html с содержимым test. Вывод ls -lZ показывает тип httpd_sys_content_t. Браузер отображает файл по HTTP. См. @fig-004.", "04.png", "fig-004"),
            ("Изменение типа и отказ доступа", "После chcon -t samba_share_t /var/www/html/test.html браузер возвращает 403 Forbidden. Вывод ls -Z подтверждает тип samba_share_t. См. @fig-005.", "05.png", "fig-005"),
            ("Анализ журнала", "В журнале присутствуют сообщения setroubleshoot о запрете обращения процесса httpd к test.html. Среди рекомендаций указано восстановление метки через restorecon. См. @fig-006.", "06.png", "fig-006"),
            ("Настройка порта Apache", "В файле /etc/httpd/conf/httpd.conf директива Listen имеет значение 81. Этот фрагмент конфигурации показан на @fig-007.", "07.png", "fig-007"),
            ("Проверка политики порта и восстановление контекста", "После перезапуска выполнены semanage port -a -t http_port_t -p tcp 81 и просмотр списка портов. Сообщение Port tcp/81 already defined, modifying instead указывает на существующее определение. Порт 81 присутствует в http_port_t. После chcon -t httpd_sys_content_t содержимое test.html снова отображается в браузере. См. @fig-008.", "08.png", "fig-008"),
        ]
        return "\n\n".join(
            f"## {title}\n\n{text}\n\n![{title}](image/{image}){{#{label} width=90%}}"
            for title, text, image, label in steps
        )
    result = []
    for part in re.split(r"(?m)^## ", source)[1:]:
        title, content = part.split("\n", 1)
        match = re.search(r"!\[([^]]*)\]\((image/[^)]+)\)\{\s*#([^\s}]+)[^}]*\}", content)
        if not match:
            continue
        title = clean_heading(title)
        text = content[:match.start()].strip()
        label = match.group(3)
        result.append(
            f"## {title}\n\n{text}\n\nРезультат показан на @{label}.\n\n"
            f"![{title}]({match.group(2)}){{#{label} width=90%}}"
        )
    expected = {2: 5, 3: 7, 4: 6, 5: 11, 7: 1, 8: 2}[lab]
    assert len(result) == expected, (lab, len(result))
    return "\n\n".join(result)


def appendix(lab, source):
    parts = []
    if lab in (2, 3):
        rows = []
        for row in source.splitlines():
            if not row.startswith("|"):
                continue
            cells = [c.strip() for c in row.strip().strip("|").split("|")]
            if len(cells) == 10 and re.search(r"\((\d{3})\)", cells[0]):
                cells[0] = re.search(r"\((\d{3})\)", cells[0]).group(1)
                cells[1] = re.search(r"\((\d{3})\)", cells[1]).group(1)
                rows.append(cells)
        assert len(rows) == 64
        legend = (
            "Обозначения: 1 — создание файла; 2 — удаление; 3 — запись; 4 — чтение; "
            "5 — вход в каталог; 6 — просмотр имён; 7 — переименование; "
            "8 — изменение обычных прав. Знак плюс означает разрешение, минус — запрет."
        )
        matrix = "| Каталог | Файл | 1 | 2 | 3 | 4 | 5 | 6 | 7 | 8 |\n|:---:|:---:|:---:|:---:|:---:|:---:|:---:|:---:|:---:|:---:|\n"
        matrix += "\n".join("| " + " | ".join(row) + " |" for row in rows)
        matrix += "\n\n: Расчётная матрица сочетаний прав {#tbl-matrix}"
        parts.append(
            "# Расчётная матрица прав\n\n" + legend + "\n\n"
            "Расчётная матрица (@tbl-matrix) является теоретической. Учтены обычные биты доступа; "
            "дополнительные ограничения ACL, SELinux, a/i и Sticky в расчёт не включены. "
            + ("Биты относятся к владельцу." if lab == 2 else "Биты относятся к группе. Изменение прав чужого файла от guest2 запрещено.")
            + "\n\n" + matrix
        )
        minimal = "| Операция | Каталог | Файл |\n|:---|:---:|:---:|\n"
        for row in source.splitlines():
            if not row.startswith("|"):
                continue
            cells = [c.strip() for c in row.strip().strip("|").split("|")]
            if len(cells) == 3 and re.search(r"\((\d{3})\)", cells[1]):
                minimal += "| " + cells[0] + " | " + re.search(r"\((\d{3})\)", cells[1]).group(1) + " | " + re.search(r"\((\d{3})\)", cells[2]).group(1) + " |\n"
        minimal += "\n: Минимальные обычные права {#tbl-minimum}"
        parts.append("# Минимальные права\n\nМинимальные комбинации приведены в @tbl-minimum. Для удаления подкаталога дополнительно требуется, чтобы он был пустым.\n\n" + minimal)
    if lab == 5:
        for name in ("simpleid.c", "simpleid2.c", "readfile.c"):
            code = (ROOT / f"labs/lab05/code/{name}").read_text()
            parts.append(f"# Листинг {name.replace('.', ' ')}\n\n{FENCE}c\n{code.rstrip()}\n{FENCE}")
    if lab in (7, 8):
        code = (ROOT / f"labs/lab{lab:02}/lab{lab:02}.py").read_text()
        parts.append(f"# Листинг программы\n\n{FENCE}python\n{code.rstrip()}\n{FENCE}")
    if lab == 4:
        commands = "ls -l file1\nlsattr file1\nsudo chattr +a file1\nsudo chattr -a file1\nsudo chattr +i file1\nsudo chattr -i file1\nlsattr file1"
        parts.append(f"# Команды управления атрибутами\n\nУстановка и снятие a/i выполняются с соответствующей привилегией.\n\n{FENCE}bash\n{commands}\n{FENCE}")
    if lab == 6:
        commands = "service httpd status\nps aux -Z | grep httpd\nsestatus -b | grep httpd\nls -lZ /var/www/html/test.html\nchcon -t samba_share_t /var/www/html/test.html\nsemanage port -l | grep http_port_t\nchcon -t httpd_sys_content_t /var/www/html/test.html"
        parts.append(f"# Команды проверки SELinux\n\nОсновные команды, рассмотренные в работе:\n\n{FENCE}bash\n{commands}\n{FENCE}\n\nПри завершении опыта возвращают исходную директиву Listen. Удаляют только созданную для опыта локальную привязку порта, если она была добавлена; существующее правило политики сохраняют. Учебный test.html удаляют после проверки.")
    parts = [re.sub(r"^# (.*)$", r"# \1 {.appendix}", part, count=1, flags=re.M) for part in parts]
    return "\n\n\\appendix\n\n" + "\n\n".join(parts) + "\n"


def flowchart(directory, labels):
    width, height = 720, 510
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 24)
    for i, label in enumerate(labels):
        y = 12 + i * 100
        draw.rounded_rectangle((25, y, 695, y + 68), radius=8, fill="#f4f4f4", outline="#555555", width=2)
        box = draw.textbbox((0, 0), label, font=font)
        draw.text(((width - (box[2] - box[0])) / 2, y + 20), label, font=font, fill="black")
        if i < len(labels) - 1:
            draw.line((360, y + 70, 360, y + 96), fill="#555555", width=3)
            draw.polygon([(360, y + 98), (353, y + 89), (367, y + 89)], fill="#555555")
    image.save(directory / "image/workflow.png")


def reference_doc():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.4)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.0)
    for name in ("Normal", "Body Text", "First Paragraph"):
        if name not in doc.styles:
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style = doc.styles[name]
        style.font.name = "DejaVu Serif"
        style.font.size = Pt(12)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.first_line_indent = Cm(0.75)
    for name, size in (("Title", 20), ("Subtitle", 14), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 12), ("Caption", 10)):
        style = doc.styles[name]
        style.font.name = "DejaVu Serif"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.keep_with_next = name != "Caption"
        if name.startswith("Heading"):
            style.paragraph_format.space_before = Pt(16)
            style.paragraph_format.space_after = Pt(8)
    for name in ("Source Code", "Verbatim Char"):
        if name in doc.styles:
            doc.styles[name].font.name = "DejaVu Sans Mono"
            doc.styles[name].font.size = Pt(9)
    footer = section.footer.paragraphs[0]
    footer.alignment = 1
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.author = "Архипов Александр Сергеевич"
    TEMPLATE.mkdir(parents=True, exist_ok=True)
    doc.save(TEMPLATE / "reference.docx")


def prepare():
    reference_doc()
    configuration = yaml.safe_load((TEMPLATE / "upstream-quarto.yml").read_text())
    configuration["format"]["pdf"].update({
        "pdf-engine": "xelatex", "cite-method": "citeproc",
        "mainfont": "DejaVu Serif", "sansfont": "DejaVu Sans", "monofont": "DejaVu Sans Mono",
        "keep-tex": False,
    })
    configuration["format"]["docx"]["reference-doc"] = "../../../templates/academic-report/reference.docx"
    configuration["crossref"].update({"fig-title": "Рисунок", "fig-prefix": "рис.", "tbl-title": "Таблица", "tbl-prefix": "табл."})
    configuration["execute"] = {"enabled": False}
    for lab in range(2, 9):
        data = SECTIONS[str(lab)]
        directory = ROOT / f"labs/lab{lab:02}/report"
        source = original(lab)
        goal = re.search(r"# Цель работы\s*\n(.*?)(?=\n# )", source, re.S).group(1).strip().replace("SELinx", "SELinux")
        if not goal.endswith("."):
            goal += "."
        meta = {
            "author": {"name": "Архипов Александр Сергеевич", "email": "1132239106@rudn.ru", "affiliation": [{"name": "Российский университет дружбы народов", "city": "Москва", "country": "Российская Федерация"}]},
            "title": f"Отчёт по лабораторной работе {lab}",
            "subtitle": data["topic"],
            "abstract": data["abstract"],
            "keywords": ["информационная безопасность", "лабораторная работа", "Linux" if lab <= 6 else "гаммирование"],
            "date": "2026-10-12", "date-format": "DD.MM.YYYY", "lang": "ru-RU",
        }
        body = (
            "# Введение\n\n"
            "Исполнитель — Архипов Александр Сергеевич, группа НБИбд-01-24. "
            "Студенческий билет 1132239106. Электронная почта: 1132239106@rudn.ru.\n\n"
            "## Актуальность\n\n" + data["relevance"] + "\n\n"
            "## Цель работы\n\n" + goal + "\n\n"
            "## Задачи\n\n" + "\n".join(f"{i}. {task}" for i, task in enumerate(data["tasks"], 1)) + "\n\n"
            "## Объект и предмет исследования\n\n" + data["object"] + "\n\n"
            "## Методы исследования\n\n" + data["methods"] + "\n\n"
            "## Схема выполнения работы\n\n"
            "Последовательность действий показана на @fig-workflow.\n\n"
            "![Схема выполнения лабораторной работы](image/workflow.png){#fig-workflow width=65%}\n\n"
            "# Теоретическая часть\n\n" + data["theory"] + "\n\n"
            "# Материалы и методы\n\n" + data["materials"] + "\n\n"
            "# Выполнение лабораторной работы\n\n" + figure_steps(lab, source) + "\n\n"
            "# Результаты\n\n" + data["results"] + "\n\n"
            "# Обсуждение результатов\n\n" + data["discussion"] + "\n\n"
            "# Заключение\n\n" + "\n".join("- " + item for item in data["conclusion"]) + "\n\n"
            "# Список литературы {.unnumbered}\n\n::: {#refs}\n:::\n"
            + appendix(lab, source)
        )
        stem = f"infosec-intro--lab{lab:02}--report"
        full = "---\n" + yaml.safe_dump(meta, allow_unicode=True, sort_keys=False, width=1000) + "---\n\n" + body
        (directory / (stem + ".md")).write_text(full)
        (directory / (stem + ".qmd")).write_text(full)
        (directory / "_quarto.yml").write_text(yaml.safe_dump(configuration, allow_unicode=True, sort_keys=False))
        (directory / "Makefile").write_text("all:\n\tquarto render\n\nclean:\n\trm -rf _output .quarto\n")
        (directory / ".gitignore").write_text(".quarto/\n_output/\n*--report.aux\n*--report.log\n*--report.lof\n*--report.lot\n*--report.tex\n*--report.toc\n*--report.out\n")
        (directory / "_resources/tex").mkdir(parents=True, exist_ok=True)
        (directory / "_resources/csl").mkdir(parents=True, exist_ok=True)
        shutil.copyfile(TEMPLATE / "preamble.tex", directory / "_resources/tex/preamble.tex")
        shutil.copyfile(TEMPLATE / "gost-r-7-0-5-2008-numeric.csl", directory / "_resources/csl/gost-r-7-0-5-2008-numeric.csl")
        keys = [key for key in dict.fromkeys(re.findall(r"@(?!fig|tbl)([a-z][a-z0-9-]+)", body)) if key in LINKS]
        bib = []
        for key in keys:
            author, title, url = LINKS[key]
            bib.append("@online{" + key + ",\n  author = {{" + author + "}},\n  title = {" + title + "},\n  url = {" + url + "},\n  urldate = {2026-10-12},\n}\n")
        (directory / "bib/cite.bib").write_text("\n".join(bib))
        flowchart(directory, data["flow"])
        print(f"Prepared report {lab:02}: {body.count('![')} figures", flush=True)


def fix_docx(path):
    """Apply academic table layout and keep pictures with their captions."""
    doc = Document(path)
    # Quarto's single-cell floating wrappers are unreliable in LibreOffice.
    # Unwrap them while retaining paragraphs, captions, bookmarks and real tables.
    for table in list(doc.tables):
        if len(table.rows) == 1 and len(table.columns) == 1:
            parent = table._tbl.getparent()
            index = parent.index(table._tbl)
            for child in list(table.cell(0, 0)._tc):
                if child.tag != qn("w:tcPr"):
                    parent.insert(index, child)
                    index += 1
            parent.remove(table._tbl)
    # Keep bookmark endpoints inside paragraphs, avoiding reader placeholder marks.
    for node in list(doc._element.body):
        if node.tag not in (qn("w:bookmarkStart"), qn("w:bookmarkEnd")):
            continue
        is_start = node.tag == qn("w:bookmarkStart")
        sibling = node.getnext() if is_start else node.getprevious()
        while sibling is not None and sibling.tag not in (qn("w:p"), qn("w:tbl")):
            sibling = sibling.getnext() if is_start else sibling.getprevious()
        if sibling is None:
            continue
        if sibling.tag == qn("w:tbl"):
            paragraphs = sibling.xpath(".//w:p")
            sibling = paragraphs[0 if is_start else -1]
        if is_start:
            sibling.insert(1 if sibling.find(qn("w:pPr")) is not None else 0, node)
        else:
            sibling.append(node)
    for style in doc.styles:
        if style._element.pPr is not None:
            for border in list(style._element.pPr.findall(qn("w:pBdr"))):
                style._element.pPr.remove(border)
    starts = {node.get(qn("w:id")) for node in doc._element.xpath(".//w:bookmarkStart")}
    ends = {node.get(qn("w:id")) for node in doc._element.xpath(".//w:bookmarkEnd")}
    orphan_ids = starts.symmetric_difference(ends)
    reference_ids = {node.get(qn("w:id")) for node in doc._element.xpath(".//w:bookmarkStart") if node.get(qn("w:name"), "").startswith("ref-") or node.get(qn("w:name")) == "refs"}
    for node in list(doc._element.xpath(".//w:bookmarkStart | .//w:bookmarkEnd")):
        if node.get(qn("w:id")) in orphan_ids | reference_ids:
            node.getparent().remove(node)
    for p in doc.paragraphs:
        if p._p.pPr is not None and p._p.pPr.pStyle is not None and p._p.pPr.pStyle.val == "Bibliography":
            p.text = p.text
    for p in doc.paragraphs:
        for border in list(p._p.get_or_add_pPr().findall(qn("w:pBdr"))):
            p._p.pPr.remove(border)
        if p.style.name in ("Title", "Subtitle"):
            for run in p.runs:
                run.font.size = Pt(20 if p.style.name == "Title" else 14)
                run.font.name = "DejaVu Serif"
                run.font.color.rgb = RGBColor(0, 0, 0)
    lab = int(re.search(r"lab(\d+)", str(path)).group(1))
    data = SECTIONS[str(lab)]
    for p in doc.paragraphs:
        if p.text in data["tasks"]:
            text = f'{data["tasks"].index(p.text) + 1}. {p.text}'
            p.text = text
            if p._p.pPr is not None and p._p.pPr.numPr is not None:
                p._p.pPr.remove(p._p.pPr.numPr)
        elif p.text in data["conclusion"]:
            p.text = "• " + p.text
            if p._p.pPr is not None and p._p.pPr.numPr is not None:
                p._p.pPr.remove(p._p.pPr.numPr)
    # Materialize a visible, linked table of contents, rather than an empty field.
    headings = []
    chapter = sub = app = 0
    in_appendix = False
    for p in doc.paragraphs:
        if p.style.name.startswith("Heading"):
            level = int(p.style.name.split()[-1])
            if level > 2:
                continue
            if level == 1:
                if p.text == "Список литературы":
                    label = p.text
                    in_appendix = True
                elif in_appendix:
                    letter = "АБВГДЕЖЗ"[app]
                    app += 1
                    p.text = "Приложение " + letter + " " + re.sub(r"^\d+[.]?\s+", "", p.text)
                    p.paragraph_format.page_break_before = True
                    number = OxmlElement("w:numPr")
                    num_id = OxmlElement("w:numId")
                    num_id.set(qn("w:val"), "0")
                    number.append(num_id)
                    old = p._p.get_or_add_pPr().numPr
                    if old is not None:
                        p._p.pPr.remove(old)
                    p._p.pPr.append(number)
                    label = p.text
                else:
                    chapter += 1
                    sub = 0
                    label = p.text if re.match(r"^\d+[.]?\s", p.text) else f"{chapter} {p.text}"
            else:
                sub += 1
                label = p.text if re.match(r"^\d+[.]\d+", p.text) else f"{chapter}.{sub} {p.text}"
            existing_ids = [int(x.get(qn("w:id"))) for x in doc._element.xpath(".//w:bookmarkStart")]
            bid = max(existing_ids, default=0) + 1
            anchor = "_academic_" + str(bid)
            start = OxmlElement("w:bookmarkStart")
            start.set(qn("w:id"), str(bid))
            start.set(qn("w:name"), anchor)
            end = OxmlElement("w:bookmarkEnd")
            end.set(qn("w:id"), str(bid))
            p._p.insert(1 if p._p.pPr is not None else 0, start)
            p._p.append(end)
            headings.append((level, label, anchor))
    for sdt in list(doc._element.body.findall(qn("w:sdt"))):
        if not sdt.xpath('.//w:instrText[contains(text(),"TOC")]', namespaces={"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}):
            continue
        title = doc.add_paragraph("Содержание", style="TOC Heading")
        title.paragraph_format.page_break_before = True
        sdt.addprevious(title._p)
        for level, label, anchor in headings:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.left_indent = Cm(0 if level == 1 else 0.5)
            paragraph.paragraph_format.line_spacing = 1
            paragraph.paragraph_format.space_after = Pt(5)
            link = OxmlElement("w:hyperlink")
            link.set(qn("w:anchor"), anchor)
            run = OxmlElement("w:r")
            text = OxmlElement("w:t")
            text.text = label
            run.append(text)
            link.append(run)
            paragraph._p.append(link)
            sdt.addprevious(paragraph._p)
        sdt.getparent().remove(sdt)
    first = next(p for p in doc.paragraphs if p.style.name == "Heading 1")
    first.paragraph_format.page_break_before = True
    # Rebuild real data tables with explicit cell widths. Quarto's inner tables
    # lack tcW and retain layout-specific XML that some office readers misplace.
    for old_table in list(doc.tables):
        values = [[cell.text for cell in row.cells] for row in old_table.rows]
        ncol = len(old_table.columns)
        table = doc.add_table(rows=len(values), cols=ncol)
        table.style = "Table Grid"
        table.autofit = False
        weights = {2: [1.2, 1.8], 3: [2, 1, 1], 4: [1, 1, 1.5, 1.5], 5: [1.5, 1, 1, 1, 1], 10: [1.5, 1.5] + [1] * 8}.get(ncol, [1] * ncol)
        widths = [Cm(16.5 * weight / sum(weights)) for weight in weights]
        for ci, width in enumerate(widths):
            table.columns[ci].width = width
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                cell.width = widths[ci]
                cell.text = values[ri][ci]
        old_table._tbl.addprevious(table._tbl)
        old_table._tbl.getparent().remove(old_table._tbl)
    for table in doc.tables:
        props = table._tbl.tblPr
        borders = props.find(qn("w:tblBorders"))
        if borders is None:
            borders = OxmlElement("w:tblBorders")
            props.append(borders)
        for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
            line = OxmlElement("w:" + name)
            for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                line.set(qn("w:" + key), value)
            borders.append(line)
        for ri, row in enumerate(table.rows):
            if ri == 0:
                repeat = OxmlElement("w:tblHeader")
                row._tr.get_or_add_trPr().append(repeat)
            for cell in row.cells:
                cell.vertical_alignment = 1
                if ri == 0:
                    fill = OxmlElement("w:shd")
                    fill.set(qn("w:fill"), "EDEDED")
                    cell._tc.get_or_add_tcPr().append(fill)
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing = 1
                    p.paragraph_format.space_before = Pt(1.5 if len(table.columns) == 10 else 3)
                    p.paragraph_format.space_after = Pt(1.5 if len(table.columns) == 10 else 3)
                    p.paragraph_format.first_line_indent = Cm(0)
                    p.paragraph_format.keep_with_next = False
                    for run in p.runs:
                        run.font.size = Pt(10)
                        run.font.color.rgb = RGBColor(0, 0, 0)
                        run.bold = ri == 0
                        run.font.name = "DejaVu Serif"
    for p in doc.paragraphs:
        if p._p.pPr is not None and p._p.pPr.pStyle is not None and p._p.pPr.pStyle.val == "SourceCode":
            p.paragraph_format.line_spacing = 1
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_with_next = False
            for run in p.runs:
                run.font.name = "DejaVu Sans Mono"
                run.font.size = Pt(9)
        if p._p.xpath(".//w:drawing"):
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.first_line_indent = Cm(0)
        if p.style.name == "Caption":
            p.paragraph_format.keep_with_next = False
            p.paragraph_format.first_line_indent = Cm(0)
        if p.text == "Список литературы" and p.style.name == "Heading 1":
            p.paragraph_format.page_break_before = True
    starts = {node.get(qn("w:id")) for node in doc._element.xpath(".//w:bookmarkStart")}
    ends = {node.get(qn("w:id")) for node in doc._element.xpath(".//w:bookmarkEnd")}
    for node in list(doc._element.xpath(".//w:bookmarkStart | .//w:bookmarkEnd")):
        if node.get(qn("w:id")) in starts.symmetric_difference(ends):
            node.getparent().remove(node)
    assert all(len(table.columns) > 1 for table in doc.tables)
    doc.core_properties.author = "Архипов Александр Сергеевич"
    doc.save(path)


def build(quarto, labs):
    for lab in labs:
        directory = ROOT / f"labs/lab{lab:02}/report"
        stem = f"infosec-intro--lab{lab:02}--report"
        for fmt in ("docx", "pdf"):
            command = [quarto, "render", stem + ".md", "--to", fmt, "--output", stem + "." + fmt]
            result = subprocess.run(command, cwd=directory, capture_output=True, text=True)
            log = ROOT / f"../tmp/template-build-lab{lab:02}-{fmt}.log"
            log.write_text(result.stdout + "\n" + result.stderr)
            if result.returncode:
                print(result.stderr[-6000:], flush=True)
                raise RuntimeError(f"Report {lab} {fmt} failed; see {log}")
            shutil.copyfile(directory / "_output" / (stem + "." + fmt), directory / (stem + "." + fmt))
            if fmt == "docx":
                fix_docx(directory / (stem + ".docx"))
            print(f"Built report {lab:02} {fmt}", flush=True)


def sources_and_manifest():
    manifest = json.loads((ROOT / "releases/labs.json").read_text())
    for release in manifest["releases"]:
        lab = release["lab"]
        folder = ROOT / f"labs/lab{lab:02}"
        archive = folder / f"LR{lab:02}_sources.zip"
        with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as output:
            for path in sorted(folder.rglob("*")):
                if not path.is_file() or path == archive:
                    continue
                if any(part in (".quarto", "_output", "__pycache__") for part in path.relative_to(folder).parts):
                    continue
                if path.suffix in (".docx", ".pdf", ".html", ".zip"):
                    continue
                if path.name.startswith("infosec-intro--lab") and path.suffix in (".aux", ".log", ".lof", ".lot", ".tex", ".toc", ".out"):
                    continue
                output.write(path, str(path.relative_to(ROOT)))
            for path in sorted(TEMPLATE.rglob("*")):
                if path.is_file():
                    output.write(path, str(path.relative_to(ROOT)))
            for name in ("report_sections.json", "reformat_reports.py"):
                output.write(ROOT / "releases" / name, "releases/" + name)
        release["tag"] = f"lab{lab:02}-v1.1.0"
        release["notes"] = (
            f"Отдельные файлы лабораторной работы №{lab}.\n\n"
            "Отчёт оформлен по шаблону yamadharma/academic-laboratory-report-template: "
            "введение, теория, материалы и методы, выполнение, результаты, обсуждение, "
            "заключение, список литературы и приложения.\n\n"
            "Отчёт: Markdown, DOCX, PDF. Презентация: Markdown, HTML, PDF. "
            "Архив исходников включает Markdown, Quarto, скриншоты, исходный код и ресурсы оформления.\n"
        )
        for asset in release["assets"]:
            path = ROOT / asset["path"]
            asset["sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
            asset["size"] = path.stat().st_size
    (ROOT / "releases/labs.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--prepare", action="store_true")
    parser.add_argument("--quarto", default="quarto")
    parser.add_argument("--labs", type=int, nargs="+", default=list(range(2, 9)))
    parser.add_argument("--package", action="store_true")
    args = parser.parse_args()
    if args.prepare:
        prepare()
    elif args.package:
        sources_and_manifest()
    else:
        build(args.quarto, args.labs)
